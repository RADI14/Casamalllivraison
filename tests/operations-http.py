"""Read-only mock carrier checks; all order writes stay in a temporary SQLite database."""
import os,subprocess,tempfile,time,urllib.request,urllib.error,json,secrets,sqlite3,uuid,io
from datetime import datetime,timezone,timedelta
from openpyxl import load_workbook
from pathlib import Path
root=str(Path(__file__).resolve().parent.parent)
with tempfile.TemporaryDirectory() as temp:
 password=secrets.token_urlsafe(24);dbfile=temp+'/app.sqlite';preload=temp+'/atlas.mjs'
 Path(preload).write_text("const original=globalThis.fetch;globalThis.fetch=(url,options={})=>{if(String(url).endsWith('/stocks'))return Promise.resolve(Response.json({success:true,data:[{stockId:'stock-test',sku:'TEST',name:'Stock product',available:100,frozen:false}]}));if(String(url).startsWith('https://api.atlaslivraison.com/')&&options.method==='POST'&&String(url).endsWith('/colis'))return Promise.resolve(Response.json({success:true,data:{code:'TEST-'+options.headers['Idempotency-Key'],cfees:35}}));if(String(url).startsWith('https://api.atlaslivraison.com/'))throw Error('Unexpected carrier call');return original(url,options)}")

 env={**os.environ,'SQLITE_PATH':dbfile,'OWNER_EMAIL':'owner@test.com','ADMIN_PASSWORD':password,'ATLAS_API_KEY':'mock-api-key','PORT':'3192','HOSTNAME':'127.0.0.1','NODE_ENV':'production','APP_URL':'http://127.0.0.1:3192'}
 log=open('/tmp/operations-http.log','w');proc=subprocess.Popen(['node','--import',preload,'.next/standalone/server.js'],cwd=root,env=env,stdout=log,stderr=log);base='http://127.0.0.1:3192'
 def req(path,body=None,cookie='',origin=base,raw=False):
  headers={'Content-Type':'application/json','Origin':origin}
  if cookie:headers['Cookie']=cookie
  request=urllib.request.Request(base+path,headers=headers,data=json.dumps(body).encode() if body is not None else None)
  try:r=urllib.request.urlopen(request);return r.status,r.headers,(r.read() if raw else json.loads(r.read()))
  except urllib.error.HTTPError as e:return e.code,e.headers,(e.read() if raw else json.loads(e.read()))
 try:
  for _ in range(100):
   try:
    if req('/api/health')[0]==200:break
   except Exception:pass
   if proc.poll() is not None:raise RuntimeError('server exited')
   time.sleep(.2)
  else:raise RuntimeError('health timed out')
  status,h,_=req('/api/auth/login',{'email':'owner@test.com','password':password});assert status==200;owner=h['Set-Cookie'].split(';')[0]
  for path in ['/api/operations','/api/operations/statement?email=a@test.com&format=xlsx']:
   assert req(path)[0]==403
  agentpass=secrets.token_urlsafe(24)
  assert req('/api/team',{'action':'create','name':'Agent','email':'a@test.com','password':agentpass},owner)[0]==200
  status,h,_=req('/api/auth/login',{'email':'a@test.com','password':agentpass});assert status==200;agent=h['Set-Cookie'].split(';')[0]
  assert req('/api/operations',cookie=agent)[0]==403
  assert req('/api/operations/statement?email=a@test.com&format=print',cookie=agent)[0]==403
  now=datetime.now(timezone.utc).isoformat();today=req('/api/operations',cookie=owner)[2]['period']['to']
  fixture=sqlite3.connect(dbfile);fixture.create_function('audit_actor',0,lambda:'test-fixture')
  fixture.execute("INSERT INTO products(id,sku,name,stock,cost_cents,sale_cents,created_at) VALUES('p','TEST','Stock product',100,2000,19900,?)",(now,))
  payload={'fullname':'=Client <script>','phone':'0612345678','phone2':'','address':'10 rue des fleurs Kenitra','cityName':'Kenitra','cityId':'city','storeId':'','product':'Stock product','quantity':2,'price':200,'note':'','openpackage':False,'change':False,'allowtry':False,'atlasSku':'TEST','stockItems':[{'stockId':'stock-test','quantity':2}]}
  fixture.execute('INSERT INTO orders(id,payload,status,tracking,created_by,created_email,commission_rate,created_at,updated_at,atlas_state,product_id) VALUES(?,?,?,?,?,?,?,?,?,?,?)',('seed',json.dumps(payload),'sent','SEED','Agent','a@test.com',10,now,now,17,'p'));fixture.commit()
  status,headers,report=req('/api/operations',cookie=owner);assert status==200 and 'no-store' in headers['Cache-Control']
  assert report['alerts']['missingPurchase']['count']==1 and report['alerts']['missingDelivery']['count']==1 and len(report['alerts']['missingAds'])==1
  assert req('/api/operations?from=2026-02-30',cookie=owner)[0]==400
  assert req('/api/operations/statement?email=unknown',cookie=owner)[0]==404
  status,_,statement=req('/api/operations/statement?email=a@test.com',cookie=owner);assert status==200 and statement['dueCents']==1000 and len(statement['rows'])==1
  status,headers,excel=req('/api/operations/statement?email=a@test.com&format=xlsx',cookie=owner,raw=True);assert status==200 and excel[:2]==b'PK' and 'attachment' in headers['Content-Disposition']
  book=load_workbook(io.BytesIO(excel),data_only=False);assert book.sheetnames==['Synthèse','Colis','Paiements','Charges'];assert book['Colis']['C2'].value=='=Client <script>' and book['Colis']['C2'].data_type=='s';assert book['Colis']['G2'].value==200 and book['Colis']['H2'].value is None
  status,_,html=req('/api/operations/statement?email=a@test.com&format=print',cookie=owner,raw=True);assert status==200 and b'window.print()' in html and b'Client <script>' not in html
  receipt={'action':'create','id':str(uuid.uuid4()),'reference':' atlas-test ','date':today,'amount':'100','note':'Received in bank','confirmed':True}
  assert req('/api/carrier-receipts',receipt,agent)[0]==403
  assert req('/api/carrier-receipts',receipt,owner,origin='https://untrusted.test')[0]==403
  assert req('/api/carrier-receipts',{**receipt,'confirmed':False},owner)[0]==400
  assert req('/api/carrier-receipts',receipt,owner)[0]==201
  assert req('/api/carrier-receipts',receipt,owner)[0]==200
  assert req('/api/carrier-receipts',{**receipt,'id':str(uuid.uuid4())},owner)[0]==409
  assert req('/api/carrier-receipts',{**receipt,'id':str(uuid.uuid4()),'reference':'FUTURE','date':(datetime.fromisoformat(today)+timedelta(days=1)).date().isoformat()},owner)[0]==400
  report=req('/api/operations',cookie=owner)[2];assert report['carrier']['receivedCents']==10000 and report['carrier']['reconciliation'] is None
  assert req('/api/carrier-receipts',{'action':'void','id':receipt['id']},owner)[0]==200
  report=req('/api/operations',cookie=owner)[2];assert report['carrier']['receivedCents']==0 and report['carrier']['receipts'][0]['voided_at']
  # Cross-agent duplicates require an explicit current snapshot without disclosing other clients.
  draft={'action':'create','requestId':str(uuid.uuid4()),'payload':{**payload,'fullname':'New client','fulfillmentMode':'stock','atlasStockId':'stock-test'}}
  status,_,warning=req('/api/orders',draft,owner);assert status==409 and warning['code']=='duplicate_order' and warning['count']==1
  assert set(warning)=={'error','code','duplicateConfirmation','count'}
  assert fixture.execute('SELECT COUNT(*) FROM orders').fetchone()[0]==1
  fixture.execute('INSERT INTO orders(id,payload,status,tracking,created_by,created_email,commission_rate,created_at,updated_at,atlas_state) VALUES(?,?,?,?,?,?,?,?,?,?)',('second',json.dumps(payload),'sent','SECOND','Agent','a@test.com',10,now,now,17));fixture.commit()
  confirmed={**draft,'allowDuplicate':True,'duplicateConfirmation':warning['duplicateConfirmation']}
  status,_,stale=req('/api/orders',confirmed,owner);assert status==409 and stale['count']==2
  confirmed['duplicateConfirmation']=stale['duplicateConfirmation'];status,_,created=req('/api/orders',confirmed,owner);assert status==201 and created['tracking']
  assert req('/api/orders',confirmed,owner)[0]==200
  assert fixture.execute('SELECT COUNT(*) FROM orders').fetchone()[0]==3
  fixture.close()
  print('PASS: admin controls, alerts, real Excel and print export, receipt validation/idempotence/void and duplicate confirmation with stale-snapshot rejection')
 finally:
  proc.terminate();proc.wait(timeout=10);log.close()
